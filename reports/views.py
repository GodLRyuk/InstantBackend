from datetime import timedelta
from django.db.models import Sum, Count, F, DecimalField
from django.db.models.functions import TruncDate, TruncMonth
from django.utils.dateparse import parse_date
from rest_framework.views import APIView
from rest_framework.response import Response
from orders.models import Order, OrderItem
from stock.models import Inventory, StockBatch
from products.models import Product
from rest_framework.permissions import IsAuthenticated
import openpyxl
from openpyxl.utils import get_column_letter
from django.http import HttpResponse


class SalesTrendAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not request.user.is_staff:
            return Response({"error": "Only admin can view reports"}, status=403)

        start_date = parse_date(request.query_params.get('start_date', ''))
        end_date = parse_date(request.query_params.get('end_date', ''))
        group_by = request.query_params.get('group_by', 'day')

        qs = Order.objects.filter(payment_status='PAID')
        if start_date:
            qs = qs.filter(created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(created_at__date__lte=end_date)

        trunc_fn = TruncMonth if group_by == 'month' else TruncDate
        raw_data = (
            qs.annotate(period=trunc_fn('created_at'))
              .values('period')
              .annotate(revenue=Sum('total_amount'), order_count=Count('id'))
              .order_by('period')
        )

        # Build a lookup of actual results
        results_by_date = {
            row["period"].date() if hasattr(row["period"], "date") else row["period"]: row
            for row in raw_data
        }

        # Fill every day in the range, defaulting to 0
        output = []
        if start_date and end_date and group_by == 'day':
            current = start_date
            while current <= end_date:
                row = results_by_date.get(current)
                output.append({
                    "period": current.isoformat(),
                    "revenue": float(row["revenue"]) if row else 0.0,
                    "order_count": row["order_count"] if row else 0,
                })
                current += timedelta(days=1)
        else:
            # fallback for month grouping or missing dates — just return raw data
            output = [
                {
                    "period": row["period"].isoformat(),
                    "revenue": float(row["revenue"] or 0),
                    "order_count": row["order_count"],
                }
                for row in raw_data
            ]

        return Response(output)

class TopProductsAPIView(APIView):
    permission_classes = [IsAuthenticated]
    """
    GET /api/reports/top-products/?limit=10
    Returns best sellers by quantity and revenue — for a bar chart.
    """
    def get(self, request):
        limit = int(request.query_params.get('limit', 10))

        data = (
            OrderItem.objects
            .filter(order__payment_status='PAID')
            .values('product__id', 'product__name')
            .annotate(
                total_quantity=Sum('quantity'),
                total_revenue=Sum(F('price') * F('quantity'), output_field=DecimalField()),
            )
            .order_by('-total_revenue')[:limit]
        )

        return Response([
            {
                "product_id": row["product__id"],
                "product_name": row["product__name"],
                "quantity_sold": row["total_quantity"],
                "revenue": float(row["total_revenue"] or 0),
            }
            for row in data
        ])


class OrderStatusBreakdownAPIView(APIView):
    permission_classes = [IsAuthenticated]
    """
    GET /api/reports/order-status/
    For a pie chart of order lifecycle distribution.
    """
    def get(self, request):
        data = (
            Order.objects.values('order_status')
            .annotate(count=Count('id'))
            .order_by('-count')
        )
        return Response(list(data))


class InventorySummaryAPIView(APIView):
    permission_classes = [IsAuthenticated]
    """
    GET /api/reports/inventory-summary/
    KPI numbers + low-stock list for the dashboard.
    """
    def get(self, request):
        inventories = Inventory.objects.select_related('product')

        total_stock_value = 0
        low_stock_items = []

        for inv in inventories:
            batch_avg_price = (
                StockBatch.objects.filter(product=inv.product)
                .aggregate(avg_price=Sum('purchase_price') / Count('id'))['avg_price']
                or 0
            )
            total_stock_value += float(inv.total_stock) * float(batch_avg_price)

            if inv.is_low_stock():
                low_stock_items.append({
                    "product_id": inv.product.id,
                    "product_name": inv.product.name,
                    "total_stock": inv.total_stock,
                    "threshold": inv.low_stock_threshold,
                })

        return Response({
            "total_stock_value": round(total_stock_value, 2),
            "total_products": inventories.count(),
            "low_stock_count": len(low_stock_items),
            "low_stock_items": low_stock_items,
        })
class SalesReportExcelAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not request.user.is_staff:
            return Response({"error": "Only admin can export reports"}, status=403)

        start_date = parse_date(request.query_params.get('start_date', ''))
        end_date = parse_date(request.query_params.get('end_date', ''))

        qs = Order.objects.filter(payment_status='PAID').select_related(
            'user', 'coupon'
        ).prefetch_related('items__product', 'items__batch')

        if start_date:
            qs = qs.filter(created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(created_at__date__lte=end_date)

        qs = qs.order_by('-created_at')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        headers = [
            "Order ID", "Date", "Customer Name", "Customer Email", "Customer Phone",
            "Product Name", "Batch No", "Quantity", "Unit Price", "Item Total",
            "Order Discount", "Delivery Fee", "Coupon Code",
            "Order Total", "Order Status", "Payment Method",
        ]
        ws.append(headers)

        for order in qs:
            customer = order.user
            customer_name = f"{customer.first_name or ''} {customer.last_name or ''}".strip() or customer.username
            coupon_code = order.coupon.code if order.coupon else ""

            items = order.items.all()

            if not items:
                # Order with no items — still show one row for visibility
                ws.append([
                    order.id, order.created_at.strftime('%Y-%m-%d %H:%M'),
                    customer_name, customer.email, getattr(customer, 'phone', ''),
                    "", "", "", "", "",
                    float(order.discount_amount), float(order.delivery_fee), coupon_code,
                    float(order.total_amount), order.order_status, order.payment_method,
                ])
                continue

            for item in items:
                ws.append([
                    order.id,
                    order.created_at.strftime('%Y-%m-%d %H:%M'),
                    customer_name,
                    customer.email,
                    getattr(customer, 'phone', ''),
                    item.product.name,
                    item.batch.batch_no if item.batch else "",
                    item.quantity,
                    float(item.price),
                    float(item.total_price),
                    float(order.discount_amount),
                    float(order.delivery_fee),
                    coupon_code,
                    float(order.total_amount),
                    order.order_status,
                    order.payment_method,
                ])

        for i, _ in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 18

        # Bold header row
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="sales_report.xlsx"'
        wb.save(response)
        return response